"""PowerPoint (.pptx) parser. """

from pathlib import Path
from parsers.base import ParsedDocument, normalize_source_file


def parse_pptx(file_path: Path) -> ParsedDocument:
    """Extract text content from PowerPoint slides."""
    from pptx import Presentation

    file_path = Path(file_path)
    if not file_path.exists():
        raise FileNotFoundError(f"PPTX not found: {file_path}")

    prs = Presentation(file_path)
    texts = []
    for i, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(f"[Slide {i+1}] {shape.text.strip()}")
    content = "\n\n".join(texts) if texts else ""
    metadata = {
        "format": "pptx",
        "slide_count": len(prs.slides),
    }

    return ParsedDocument(
        content=content,
        source_file=normalize_source_file(file_path),
        metadata=metadata,
    )