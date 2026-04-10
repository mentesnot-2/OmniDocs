"""
XLSX (Microsoft Excel) text extraction using openpyxl.
"""
from pathlib import Path
from parsers.base import ParsedDocument, normalize_source_file

def parse_xlsx(file_path:Path) -> ParsedDocument:
    """Extract text from all sheets in an Excel file."""
    from openpyxl import load_workbook

    file_path = Path(file_path)
    if not file_path.exists():
        raise FileNotFoundError(f"XLSX not found: {file_path}")

    wb = load_workbook(str(file_path),read_only=True,data_only=True)
    parts = []

    for sheet in wb.worksheets:
        sheet_lines = [f"Sheet: {sheet.title}"]
        for row in sheet.iter_rows(values_only=True):
            row_values = [str(c) if c is not None else "" for c in row]
            if any(v.strip() for v in row_values):
                sheet_lines.append(" | ".join(v.strip() for v in row_values))
        if len(sheet_lines) > 1:
            parts.append("\n".join(sheet_lines))
    full_text = "\n\n".join(parts)
    metadata = {
        "format":"xlsx",
        "sheet_count": len(wb.worksheets),
    }
    wb.close()
    return ParsedDocument(
        content=full_text,
        source_file=normalize_source_file(file_path),
        metadata=metadata,
    )