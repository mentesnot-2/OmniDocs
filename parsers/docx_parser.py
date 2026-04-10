"""
DOCX (Microsoft Word) text extraction using python-docx.
"""

from pathlib import Path
from parsers.base import ParsedDocument, normalize_source_file

def parse_docx(file_path:Path) -> ParsedDocument:
    """
    Extract text from .docx file.
    Iterates through paragraphs and tables to capture all content.
    """
    from docx import Document

    file_path = Path(file_path)
    if not file_path.exists():
        raise FileNotFoundError(f"DOCX not found: {file_path}")
    doc = Document(file_path)
    parts = []

    for para in doc.paragraphs:
        text = para.text
        if text:
            parts.append(text)

    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                if cell.text.strip():
                    row_text.append(cell.text.strip())
            if row_text:
                parts.append(" | ".join(row_text))
    full_text = "\n\n".join(parts)
    metadata = {
        "format":"docx",
        "paragraph_count": len([p for p in doc.paragraphs if p.text.strip()]),
        "table_count": len(doc.tables),
    }

    return ParsedDocument(
        content=full_text,
        source_file=normalize_source_file(file_path),
        metadata=metadata,
    )